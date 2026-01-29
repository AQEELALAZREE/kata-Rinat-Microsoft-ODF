#!/usr/bin/env python3
"""
Generate Excel test files from ODF formula specifications.

This script reads ODF formula specs and creates Excel files with those formulas
for validation against real Excel behavior.

Usage:
    python generate_excel_tests.py specs/odf-formulas/ -o test-cases/generated/
"""

import argparse
import xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import Workbook


class ODFSpecParser:
    """Parse ODF formula specifications"""
    
    def __init__(self, specs_dir):
        self.specs_dir = Path(specs_dir)
        self.formulas = []
    
    def parse_all(self):
        """Parse all ODF spec files"""
        for spec_file in self.specs_dir.glob('**/*.xml'):
            self.parse_file(spec_file)
        return self.formulas
    
    def parse_file(self, spec_file):
        """Parse single ODF spec file"""
        try:
            tree = ET.parse(spec_file)
            root = tree.getroot()
            
            # Find all formula elements
            for formula_elem in root.findall('.//formula'):
                formula_data = {
                    'name': formula_elem.get('name', 'unnamed'),
                    'expression': formula_elem.findtext('expression', ''),
                    'expected': self.parse_expected(formula_elem.find('expected')),
                    'description': formula_elem.findtext('description', ''),
                    'category': formula_elem.get('category', 'general'),
                    'source_file': spec_file.name
                }
                self.formulas.append(formula_data)
        except ET.ParseError as e:
            print(f"Warning: Could not parse {spec_file}: {e}")
    
    def parse_expected(self, expected_elem):
        """Parse expected value with type information"""
        if expected_elem is None:
            return None
        
        value_type = expected_elem.get('type', 'auto')
        value_text = expected_elem.text
        
        if value_type == 'number':
            return float(value_text)
        elif value_type == 'string':
            return value_text
        elif value_type == 'boolean':
            return value_text.upper() == 'TRUE'
        elif value_type == 'error':
            return value_text  # e.g., "#DIV/0!"
        else:
            # Auto-detect
            try:
                return float(value_text)
            except (ValueError, TypeError):
                return value_text


class ExcelTestGenerator:
    """Generate Excel test files from formula specs"""
    
    def __init__(self, formulas):
        self.formulas = formulas
    
    def generate(self, output_file):
        """Generate Excel file with all formulas"""
        wb = Workbook()
        
        # Group formulas by category
        categories = {}
        for formula in self.formulas:
            category = formula['category']
            if category not in categories:
                categories[category] = []
            categories[category].append(formula)
        
        # Create sheet for each category
        for i, (category, formulas) in enumerate(categories.items()):
            if i == 0:
                ws = wb.active
                ws.title = category[:31]  # Excel sheet name limit
            else:
                ws = wb.create_sheet(category[:31])
            
            # Add headers
            ws['A1'] = 'Name'
            ws['B1'] = 'Formula'
            ws['C1'] = 'Result'
            ws['D1'] = 'Expected'
            ws['E1'] = 'Match'
            ws['F1'] = 'Description'
            
            # Add formulas
            for row, formula in enumerate(formulas, start=2):
                ws[f'A{row}'] = formula['name']
                ws[f'B{row}'] = formula['expression']
                
                # Put formula in C column (Excel will calculate it)
                if formula['expression'].startswith('='):
                    ws[f'C{row}'] = formula['expression']
                else:
                    ws[f'C{row}'] = '=' + formula['expression']
                
                # Expected value
                ws[f'D{row}'] = formula['expected']
                
                # Match formula (compares result with expected)
                ws[f'E{row}'] = f'=IF(C{row}=D{row},"✓","✗")'
                
                # Description
                ws[f'F{row}'] = formula['description']
        
        # Save workbook
        output_path = Path(output_file)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        
        print(f"Generated Excel file: {output_path}")
        print(f"Total formulas: {len(self.formulas)}")
        print(f"Categories: {len(categories)}")
        
        return output_path


def main():
    parser = argparse.ArgumentParser(
        description='Generate Excel test files from ODF specs'
    )
    parser.add_argument('specs_dir', help='Directory with ODF spec files')
    parser.add_argument('-o', '--output', default='test-cases/generated/formulas.xlsx',
                       help='Output Excel file')
    
    args = parser.parse_args()
    
    # Parse ODF specs
    print(f"Parsing ODF specs from {args.specs_dir}...")
    spec_parser = ODFSpecParser(args.specs_dir)
    formulas = spec_parser.parse_all()
    
    if not formulas:
        print("Warning: No formulas found in specs")
        return
    
    # Generate Excel file
    generator = ExcelTestGenerator(formulas)
    generator.generate(args.output)
    
    print("\nNext step: Open the Excel file and let Excel calculate the formulas")
    print(f"Then run: python scripts/validate_with_excel.py {args.output}")


if __name__ == '__main__':
    main()
