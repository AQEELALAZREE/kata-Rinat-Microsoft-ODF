import json
import argparse
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec_json", help="Path to JSON spec file")
    ap.add_argument("-o", "--output", required=True, help="Output .xlsx path")
    args = ap.parse_args()

    spec = json.loads(Path(args.spec_json).read_text(encoding="utf-8"))

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for sheet in spec["sheets"]:
        ws = wb.create_sheet(sheet["name"][:31])

        # First: write all provided cells (values/formulas)
        for addr, cell in sheet["cells"].items():
            if cell.get("formula"):
                ws[addr].value = cell["formula"]
            else:
                ws[addr].value = cell.get("value")

        # Second: ensure there is an Expected column next to Result (B -> C)
        # If the sheet uses B1="Result", we set C1="Expected"
        if ws["B1"].value and str(ws["B1"].value).strip().lower() == "result":
            ws["C1"].value = "Expected"

        # Third: for each formula cell, write expected (value or error) into the cell to the right
        # This assumes the "Result" formulas are in column B (like your specs: B2..B15).
        for addr, cell in sheet["cells"].items():
            if not cell.get("formula"):
                continue

            col_letters, row = coordinate_from_string(addr)
            col_idx = column_index_from_string(col_letters)

            # Only auto-write expected next to column B formulas (Result column)
            if col_idx != 2:  # 2 == column B
                continue

            expected = None
            if "value" in cell:
                expected = cell["value"]
            elif "error" in cell:
                expected = cell["error"]

            if expected is None:
                continue

            expected_addr = f"{get_column_letter(col_idx + 1)}{row}"  # C{row}
            ws[expected_addr].value = expected

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote: {out}")

if __name__ == "__main__":
    main()
